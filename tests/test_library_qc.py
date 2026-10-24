"""Library QC table (ytab.qc.library_qc): flag rules, count-distribution plot data and a full build on a synthetic project."""
import csv
import json
from pathlib import Path

import numpy as np
import pytest

from ytab.qc import library_qc as qc

HEADER = ["Chromosome", "Source", "Up feature type", "Up feature name", "Up gene name", "Up feature dist", "Down feature type",
          "Down feature name", "Down gene name", "Down feature dist", "IntergenicType", "Hit position", "Hit count"]


def test_skewness_rule_of_thumb():
    flags = qc.flag({"a":10., "b":40., "c":60., "d":None}, dict.fromkeys("abcd", "parent"), "skewness", "high", "rule_of_thumb")
    assert [flags[k][0] for k in "abcd"] == ["green", "yellow", "red", "grey"]


def test_project_relative_rule_is_within_role_and_directional():
    values = {"p1":100., "p2":102., "p3":98., "p4":101., "p5":40., "t1":40., "t2":41.}
    roles = {**dict.fromkeys(["p1", "p2", "p3", "p4", "p5"], "parent"), "t1":"treated", "t2":"treated"}
    low = qc.flag(values, roles, "unique_sites", "low", "relative")
    assert low["p5"][0] == "red" and all(low[k][0] == "green" for k in ("p1", "p2", "p3", "p4"))
    assert low["t1"][0] == "grey" and "fewer than 4 treated" in low["t1"][1]  # treated not compared with parents
    high = qc.flag(values, roles, "unique_sites", "high", "relative")
    assert high["p5"][0] == "green"  # only the bad direction is flagged


def test_spread_floor_ignores_trivial_differences():
    values = {"a":24.39, "b":24.97, "c":25.00, "d":24.95}
    assert all(f[0] == "green" for f in qc.flag(values, dict.fromkeys(values, "parent"), "x", "both", "relative").values())


def test_descriptive_metrics_are_not_flagged():
    assert {f[0] for f in qc.flag({"a":1., "b":9.}, dict.fromkeys("ab", "parent"), "kurtosis", None, None).values()} == {"grey"}


def test_distribution_plot_data():
    counts = np.array([1, 1, 2, 3, 5, 8, 13, 400], dtype=float)
    data = qc.plot_data("s", counts, n_points=20)
    assert sum(r["sites"] for r in data["histogram"]) == len(counts)
    assert data["ranked"][0]["rank"] == 1 and data["ranked"][0]["reads"] == 400
    qq = data["qq"]; assert all(a["observed"] <= b["observed"] and a["geometric"] <= b["geometric"] for a, b in zip(qq, qq[1:]))
    lorenz = data["lorenz"]   # the top site holds 400 of 433 reads; all sites hold all reads
    assert lorenz[0]["fraction_sites"] == 1/8 and abs(lorenz[0]["share_reads"] - 400/433) < 1e-6 and lorenz[-1]["share_reads"] == 1


def _hits(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="") as handle:
        w = csv.writer(handle, delimiter="\t", lineterminator="\n"); w.writerow(HEADER)
        for chrom, strand, pos, n in rows:w.writerow([chrom, strand, "gene", "g", "G", "1", "gene", "h", "H", "2", "W-W", pos, n])


def test_build_on_a_synthetic_project(tmp_path):
    root = tmp_path; project = root/"output/projects/P"; (project/"config").mkdir(parents=True)
    (root/"ref.fa").write_text(">c1\n" + "A"*10000 + "\n")
    rng = np.random.default_rng(1); names = [f"p{i}" for i in range(5)]
    for i, name in enumerate(names):
        sites = rng.choice(np.arange(1, 10000), 300 if i < 4 else 60, replace=False)
        _hits(project/"create_hit_file"/name/f"{name}_hits.txt", [("c1", "W" if k % 2 else "C", int(p), int(rng.integers(1, 50))) for k, p in enumerate(sites)])
    with open(project/"config/sample_sheet.csv", "w", newline="") as handle:
        w = csv.writer(handle); w.writerow(["sample", "library_role", "include"])
        for name in names:w.writerow([name, "parent", "True"])
    (project/"config/project.yaml").write_text(f"project_id: P\nrepo_root: {root}\noutput_project_dir: output/projects/P\n"
                                               f"sample_sheet: output/projects/P/config/sample_sheet.csv\nreference:\n  fasta: ref.fa\n")
    manifest = qc.build(project/"config/project.yaml", tmp_path/"out")
    assert manifest["libraries"] == 5 and manifest["genome_bp"] == 10000 and not manifest["libraries_without_hits"]
    table = {r["sample"]:r for r in csv.DictReader(open(tmp_path/"out/qc_table.tsv"), delimiter="\t")}
    assert float(table["p0"]["insertions_per_kb"]) == pytest.approx(30.0)
    flags = {(r["sample"], r["metric"]):r for r in csv.DictReader(open(tmp_path/"out/qc_flags.tsv"), delimiter="\t")}
    assert flags[("p4", "unique_sites")]["flag"] == "red" and flags[("p0", "unique_sites")]["flag"] == "green"
    assert flags[("p4", "input_reads")]["flag"] == "grey"  # no cleanup record: not available, not failed
    assert {r["metric"] for r in csv.DictReader(open(tmp_path/"out/metrics.tsv"), delimiter="\t")} == {m[0] for m in qc.METRICS}
    from openpyxl import load_workbook
    sheet = load_workbook(tmp_path/"out/qc_table.xlsx")["QC"]
    column = [c.value for c in sheet[1]].index("Insertion sites")+1
    cell = sheet.cell(6, column)
    assert cell.fill.fgColor.rgb.endswith("F4C7C3") and "robust z" in cell.comment.text
    with pytest.raises(ValueError, match="new, empty"):qc.build(project/"config/project.yaml", tmp_path/"out")
    assert json.loads((tmp_path/"out/manifest.json").read_text())["rules"]["comparison"].startswith("within role")
